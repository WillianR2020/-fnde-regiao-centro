"""Miniatura e hyperlink na coluna PRINT.

A conta que decide tudo: 2.160 prints a 400 KB dariam um .xlsx de ~860 MB,
que não abre. Comprimido a ~50 KB e dividido por categoria, a maior
planilha fica em ~60 MB — abre em qualquer máquina.

Por isso o PNG original fica intocado em data/output/prints/ como prova, e
o que entra na planilha é um JPEG reduzido em data/output/miniaturas/.
"""

from __future__ import annotations

import os
from pathlib import Path

from src.core.log import obter

log = obter(__name__)

# Formato da entrega do grupo (EQUIPAMENTO.xlsx, rodada Sul): cada print
# aparece num quadro de 376 x 400 px na coluna PRINT, numa linha de 300
# pontos, e a miniatura é o TOPO da página (onde ficam título, preço e a
# área de compra) recortado nessa proporção, salvo em 772 px de largura.
# A página inteira continua no PNG original, aberto pelo hyperlink.
QUADRO_PX = (376, 400)
LARGURA_MINIATURA = 772
QUALIDADE = 70
ALTURA_LINHA = QUADRO_PX[1]
LARGURA_COLUNA_PRINT = 53.7

# Excel mede altura de linha em pontos e largura de coluna em caracteres.
# 1 ponto = 4/3 pixel; a largura de 1 caractere fica em ~7 px na fonte
# padrão. As duas conversões existem só para a miniatura caber na célula.
PONTOS_POR_PIXEL = 0.75
PIXELS_POR_CARACTERE = 7.0

MINIATURAS = Path("data/output/miniaturas")


def caminho_miniatura(original: Path) -> Path:
    """Espelha a estrutura de prints/ dentro de miniaturas/."""
    original = Path(original)
    try:
        relativo = original.relative_to("data/output/prints")
    except ValueError:
        relativo = Path(original.name)
    return MINIATURAS / relativo.with_suffix(".jpg")


def gerar_miniatura(origem: Path, destino: Path | None = None) -> Path:
    """Reduz para LARGURA_MINIATURA e salva JPEG.

    Nunca sobrescreve o PNG original — ele é a prova. Se a miniatura já
    existe e é mais nova que o original, não refaz: montar a entrega é
    uma operação que se roda dezenas de vezes.
    """
    from PIL import Image

    origem = Path(origem)
    destino = Path(destino) if destino else caminho_miniatura(origem)

    if not origem.exists():
        raise FileNotFoundError(f"print nao existe: {origem}")

    if (destino.exists() and destino.stat().st_mtime >= origem.stat().st_mtime
            and _no_formato(destino)):
        return destino

    destino.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(origem) as imagem:
        imagem = imagem.convert("RGB")
        largura, altura = imagem.size
        proporcao = QUADRO_PX[0] / QUADRO_PX[1]
        imagem = imagem.crop((0, 0, largura, min(altura, round(largura / proporcao))))
        imagem.thumbnail((LARGURA_MINIATURA, LARGURA_MINIATURA * 4), Image.LANCZOS)
        imagem.save(destino, "JPEG", quality=QUALIDADE, optimize=True)

    return destino


def _no_formato(miniatura: Path) -> bool:
    """A miniatura já existente foi feita no formato atual (largura e recorte)?"""
    from PIL import Image

    try:
        with Image.open(miniatura) as imagem:
            largura, altura = imagem.size
    except OSError:
        return False
    return largura <= LARGURA_MINIATURA and altura <= round(largura * QUADRO_PX[1] / QUADRO_PX[0]) + 1


def embutir(ws, linha: int, coluna: int, miniatura: Path, original: Path,
            relativo_a: Path | None = None) -> None:
    """Ancora a miniatura na célula e aponta o hyperlink para o original.

    A imagem ocupa o QUADRO_PX inteiro: a miniatura já sai recortada na
    proporção dele. A célula leva o caminho do PNG relativo à pasta da
    planilha (`prints/{CATEGORIA}/...png`), como na entrega do grupo, e
    o hyperlink aponta para o mesmo caminho: a pasta `prints/` viaja
    junto com o .xlsx e o link continua funcionando em outra máquina.
    """
    from openpyxl.drawing.image import Image as ImagemXLSX
    from openpyxl.utils import get_column_letter

    miniatura, original = Path(miniatura), Path(original)
    if not miniatura.exists():
        return

    imagem = ImagemXLSX(str(miniatura))
    proporcao = (imagem.width / imagem.height) if imagem.height else 1.0
    imagem.width, imagem.height = QUADRO_PX
    if proporcao > QUADRO_PX[0] / QUADRO_PX[1]:   # página mais baixa que o quadro
        imagem.height = max(1, int(QUADRO_PX[0] / proporcao))

    celula = f"{get_column_letter(coluna)}{linha}"
    imagem.anchor = celula
    ws.add_image(imagem)

    ws.row_dimensions[linha].height = ALTURA_LINHA * PONTOS_POR_PIXEL
    largura_atual = ws.column_dimensions[get_column_letter(coluna)].width or 0
    ws.column_dimensions[get_column_letter(coluna)].width = max(
        largura_atual, LARGURA_COLUNA_PRINT
    )

    # O hyperlink fica na própria célula PRINT, não numa coluna a mais: o
    # template do FNDE tem 15 colunas e a PRINT é a 16ª combinada. Quem
    # quiser o PNG de prova clica na célula.
    alvo = ws.cell(row=linha, column=coluna)
    try:
        if relativo_a is not None:
            caminho = Path(os.path.relpath(original.resolve(), Path(relativo_a).resolve())).as_posix()
        else:
            caminho = original.resolve().as_uri()
    except (OSError, ValueError):
        try:
            caminho = original.resolve().as_uri()
        except (OSError, ValueError):
            log.debug("nao consegui montar o hyperlink para %s", original)
            return
    alvo.value = caminho
    alvo.hyperlink = caminho
    alvo.style = "Hyperlink"
